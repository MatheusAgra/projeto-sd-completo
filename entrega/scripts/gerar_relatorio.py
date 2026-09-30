import argparse
import base64
import csv
from io import BytesIO
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MODULES = [
    'somador_1bit',
    'sm_para_c2',
    'somador_subtrator_6bit',
    'c2_para_sm',
    'negador_c2_6bit',
    'modulo2_soma_sub',
    'comparador_c2_6bit',
    'logica_5bit',
    'decodificador_operacao',
    'mux_resultado_8x6',
    'ula_core',
    'somador_subtrator_5bit',
    'bin_bcd',
    'bcd_7seg',
    'display_decimal_2digitos',
    'ula_de2_115',
]
BLUE = '1F4E78'
PALE = 'EAF1F8'
GRID = 'D9D9D9'
BLACK = '000000'


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        properties.append(margins)
    for edge, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = margins.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        properties.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '5')
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), GRID)


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement('w:tblHeader')
    marker.set(qn('w:val'), 'true')
    properties.append(marker)


def add_table(document, headers, rows, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    table.style = 'Table Grid'
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.text = str(label)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = str(value)
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[column])
            if row_index % 2:
                set_cell_shading(cells[column], PALE)
            for paragraph in cells[column].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, separate, text, end])


def prepare_styles(document):
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.footer_distance = Inches(0.35)
    add_page_number(section.footer.paragraphs[0])
    styles = document.styles
    for name, size in (('Normal', 10.5), ('Title', 27), ('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 11)):
        style = styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        if name.startswith('Heading'):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    styles['Normal'].paragraph_format.space_after = Pt(6)
    styles['Normal'].paragraph_format.line_spacing = 1.1
    styles['Title'].font.bold = True
    styles['Title'].paragraph_format.space_after = Pt(14)


def add_text(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def add_picture(document, source, caption, max_width=6.35, max_height=3.7):
    from PIL import Image

    stream = BytesIO(source) if isinstance(source, bytes) else source
    with Image.open(stream) as image:
        width, height = image.size
    scale = min(max_width / (width / 96), max_height / (height / 96))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(BytesIO(source) if isinstance(source, bytes) else str(source), width=Inches(width / 96 * scale), height=Inches(height / 96 * scale))
    add_caption(document, caption)


def add_circuit_panels(document, source, caption, max_width=6.35, panel_height=3.45):
    from PIL import Image

    stream = BytesIO(source) if isinstance(source, bytes) else source
    with Image.open(stream) as image:
        image = image.convert('RGB')
        width, height = image.size
        rows_per_panel = max(1, int(width * panel_height / max_width))
        panel_count = (height + rows_per_panel - 1) // rows_per_panel
        for panel_index, top in enumerate(range(0, height, rows_per_panel), start=1):
            if panel_index > 1:
                document.add_page_break()
            panel = image.crop((0, top, width, min(height, top + rows_per_panel)))
            stream = BytesIO()
            panel.save(stream, format='PNG')
            add_picture(document, stream.getvalue(), f'{caption} — painel {panel_index}/{panel_count}', max_width=max_width, max_height=panel_height)


def strip_inline(text):
    return re.sub(r'`([^`]+)`', r'\1', text.strip())


def add_markdown(document, text):
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('#'):
            level = min(3, len(line) - len(line.lstrip('#')))
            heading = strip_inline(line.lstrip('#').strip())
            if heading.lower() not in ('bin_bcd.bdf', 'bin_bcd.bsf', 'bcd_7seg.bdf', 'bcd_7seg.bsf', 'display_decimal_2digitos.bdf', 'display_decimal_2digitos.bsf'):
                document.add_heading(heading, level=level)
            index += 1
            continue
        if line.startswith('|') and index + 1 < len(lines) and re.match(r'^\s*\|?\s*:?-{2,}', lines[index + 1]):
            headers = [strip_inline(cell) for cell in line.strip('|').split('|')]
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append([strip_inline(cell) for cell in lines[index].strip().strip('|').split('|')])
                index += 1
            add_table(document, headers, rows)
            continue
        if line.startswith('- ') or line.startswith('* '):
            paragraph = document.add_paragraph(style='List Bullet')
            paragraph.add_run(strip_inline(line[2:]))
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.add_run(strip_inline(line))
        index += 1


def interface_rows(interface):
    rows = []
    for name, spec in interface['ports'].items():
        width = spec['width']
        signal = name if width == 1 else f'{name}[{width - 1}:0]'
        rows.append([signal, 'Entrada' if spec['direction'] == 'input' else 'Saída', str(width)])
    return rows


def add_pin_assignment_tables(document, path):
    descriptions = {
        'SW': 'Operandos em sinal e magnitude e seletor de operação',
        'LEDR': 'Espelhamento de SW e saídas fixas em zero',
        'LEDG': 'Resultado F, STATUS e saídas fixas em zero',
        'HEX0': 'Unidades de F',
        'HEX1': 'Dezenas de F',
        'HEX2': 'Unidades de B',
        'HEX3': 'Dezenas de B',
        'HEX4': 'Unidades de A',
        'HEX5': 'Dezenas de A',
        'HEX6': 'Display apagado',
        'HEX7': 'Display apagado',
    }
    segments = ('a', 'b', 'c', 'd', 'e', 'f', 'g')

    def usage(port, bit):
        if port == 'SW':
            if bit < 4:
                return f'magnitude A[{bit}]'
            if bit == 4:
                return 'sinal A'
            if bit < 9:
                return f'magnitude B[{bit - 5}]'
            if bit == 9:
                return 'sinal B'
            return f'seleção S[{bit - 10}]'
        if port == 'LEDR':
            return f'espelho SW[{bit}]' if bit < 13 else 'fixo em zero'
        if port == 'LEDG':
            return f'F[{bit}]' if bit < 6 else ('STATUS' if bit == 6 else 'fixo em zero')
        return f'segmento {segments[bit]} (ativo em zero)'

    groups = {}
    with path.open(newline='', encoding='utf-8-sig') as stream:
        for row in csv.DictReader(stream):
            match = re.fullmatch(r'(SW|LEDR|LEDG|HEX[0-7])\[(\d+)\]', row['sinal'])
            if match:
                groups.setdefault(match.group(1), []).append((int(match.group(2)), row))
    for port in ('SW', 'LEDR', 'LEDG', *(f'HEX{i}' for i in range(8))):
        rows = groups.get(port, [])
        if not rows:
            continue
        document.add_heading(f'{port} — {descriptions[port]}', level=3)
        rows.sort(key=lambda item: item[0])
        add_table(document, ['Bit', 'Sinal', 'Uso', 'Pino FPGA', 'Padrão I/O'], [
            [bit, row['sinal'], usage(port, bit), row['pino'], row['padrao_io']] for bit, row in rows
        ], font_size=8)


def get_assets(root, circuit_dir, waveform_dir, validation):
    missing = []
    if not validation.is_file():
        missing.append(str(validation))
    circuits = {}
    for module in MODULES:
        png = circuit_dir / f'{module}.png'
        svg = circuit_dir / f'{module}.svg'
        if png.is_file():
            circuits[module] = png
        elif svg.is_file():
            circuits[module] = svg
        else:
            missing.append(f'{png} ou {svg}')
        waveform = waveform_dir / f'{module}.png'
        if not waveform.is_file():
            missing.append(str(waveform))
    integration = waveform_dir / 'ula_integracao.png'
    if not integration.is_file():
        missing.append(str(integration))
    for module in MODULES:
        for suffix in ('bdf.md', 'bsf.md'):
            path = root / 'modulos' / f'{module}.{suffix}'
            if not path.is_file():
                missing.append(str(path))
    architecture = root / 'docs' / 'diagramas' / 'visao_geral.svg'
    if not architecture.is_file():
        missing.append(str(architecture))
    table_paths = [
        root / 'docs' / 'tabelas' / 'bin_bcd_tabela.csv',
        root / 'docs' / 'tabelas' / 'bcd_7seg_tabela.csv',
        root / 'docs' / 'tabelas' / 'display_decimal_2digitos_tabela.csv',
        root / 'docs' / 'tabelas' / 'bin_bcd_kmaps.svg',
        root / 'docs' / 'tabelas' / 'bcd_7seg_kmaps.svg',
        root / 'docs' / 'tabelas' / 'reducoes_bcd_displays.md',
        root / 'docs' / 'OPERACAO_010_ALTERNATIVAS.md',
        root / 'config' / 'pinagem_de2_115.csv',
        root / 'docs' / 'tabelas' / 'aritmetica' / 'somador_1bit.csv',
        root / 'docs' / 'tabelas' / 'aritmetica' / 'kmap_somador_1bit.svg',
        root / 'docs' / 'tabelas' / 'aritmetica' / 'sm_para_c2.csv',
        root / 'docs' / 'tabelas' / 'aritmetica' / 'c2_para_sm.csv',
        root / 'docs' / 'tabelas' / 'aritmetica' / 'negador_c2_6bit.csv',
        root / 'docs' / 'tabelas' / 'mapa_logica_5bit.md',
        root / 'docs' / 'tabelas' / 'tabela_decodificador_operacao.md',
    ]
    missing.extend(str(path) for path in table_paths if not path.is_file())
    return missing, integration, circuits


def rasterize_svg(node, node_modules, source):
    environment = os.environ.copy()
    environment['NODE_PATH'] = str(node_modules)
    code = "const sharp=require('sharp');sharp(process.argv[1],{density:220}).png().toBuffer().then(b=>process.stdout.write(b.toString('base64'))).catch(e=>{console.error(e);process.exit(1)})"
    result = subprocess.run([str(node), '-e', code, str(source)], env=environment, capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    return base64.b64decode(result.stdout)


def add_csv_table(document, path, headers, selector, limit=None):
    with path.open(newline='', encoding='utf-8-sig') as stream:
        rows = list(csv.DictReader(stream))
        selected = [value for row in rows if (value := selector(row)) is not None]
    if limit is not None:
        selected = selected[:limit]
    add_table(document, headers, selected, font_size=8)


def build(args):
    root = args.delivery.resolve()
    circuit_dir = args.circuit_dir.resolve()
    waveform_dir = args.waveform_dir.resolve()
    validation = args.validation.resolve()
    output = args.output.resolve()
    missing, integration, circuits = get_assets(root, circuit_dir, waveform_dir, validation)
    if missing:
        raise FileNotFoundError('Relatório aguardando insumos de validação e imagens:\n' + '\n'.join(missing))
    interfaces = json.loads((root / 'config' / 'interfaces.json').read_text(encoding='utf-8'))
    names = set(MODULES)
    if names != set(interfaces):
        raise ValueError('A lista de módulos do relatório diverge de interfaces.json')
    runtime_dependencies = Path(sys.executable).resolve().parent.parent
    node = args.node.resolve() if args.node else runtime_dependencies / 'node' / 'bin' / 'node.exe'
    node_modules = args.node_modules.resolve() if args.node_modules else runtime_dependencies / 'node' / 'node_modules'
    if not node.is_file():
        raise FileNotFoundError(f'Runtime Node.js bundled ausente: {node}')
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    prepare_styles(document)
    properties = document.core_properties
    properties.title = 'Relatório técnico da ULA'
    properties.subject = 'Implementação estrutural e validação da ULA'
    properties.author = ''
    properties.last_modified_by = ''
    properties.keywords = 'ULA; BDF; Quartus; lógica combinacional'
    title = document.add_paragraph(style='Title')
    title.add_run('Relatório técnico da ULA')
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(24)
    subtitle.add_run('Implementação estrutural em BDF e validação').font.size = Pt(16)
    add_text(document, 'Dispositivo EP4CE115F29C7')
    add_text(document, 'Integrantes: ______________________________________________')
    add_text(document, 'Professor(a): _____________________________________________')
    add_text(document, 'Disciplina: _______________________________________________')
    add_text(document, f'Data: {args.date}')
    document.add_page_break()
    document.add_heading('Resumo', level=1)
    add_text(document, 'Este relatório apresenta a implementação estrutural da ULA combinacional, seus conversores, operadores, controle, displays e integração com a placa DE2-115. A compilação final registrou zero erros e 36 avisos; os 16 módulos e a integração passaram os testbenches, e a netlist funcional nativa passou as 16.384 verificações. O fechamento do pacote e a recompilação do ZIP candidato estão detalhados em docs/PACOTE_VALIDADO.md. Não houve teste físico na placa.')
    document.add_heading('Especificação e arquitetura', level=1)
    add_text(document, 'A ULA recebe dois operandos de cinco bits em sinal e magnitude, com sinal em bit 4 e magnitude em bits 3..0. Soma e subtração produzem resultado em sinal e magnitude de seis bits. A operação 010 encaminha o complemento de B em complemento de dois de seis bits, conforme interpretação provisória registrada no projeto. Comparações usam operandos convertidos para C2; AND e XOR atuam sobre os cinco bits brutos e preservam o alinhamento acordado na integração.')
    add_picture(document, rasterize_svg(node, node_modules, root / 'docs' / 'diagramas' / 'visao_geral.svg'), 'Visão geral da arquitetura e dos sinais da DE2-115', max_width=6.35, max_height=6.4)
    document.add_heading('Interpretação da operação 010', level=2)
    add_markdown(document, (root / 'docs' / 'OPERACAO_010_ALTERNATIVAS.md').read_text(encoding='utf-8-sig'))
    add_table(document, ['Família', 'Papel'], [
        ['Conversão e aritmética', 'SM↔C2, soma/subtração de 6 bits, negação C2 e conversão para sinal e magnitude'],
        ['Operações e controle', 'Comparação signed, AND/XOR de cinco bits, decodificação one-hot e multiplexação'],
        ['BCD e displays', 'Conversão de MAG[4:0] para dezenas/unidades, decodificação ativa em zero e máscara ENABLE'],
        ['Integração', 'SW e LEDs, oito displays e sinais de status do topo ula_de2_115'],
    ])
    document.add_heading('Mapeamento físico de SW, LED e HEX', level=2)
    add_text(document, 'As tabelas abaixo mostram a função de cada grupo e cada pino FPGA. SW[4] e SW[9] são os sinais dos operandos; SW[12..10] seleciona a operação. HEX usa segmentos ativos em zero.')
    add_pin_assignment_tables(document, root / 'config' / 'pinagem_de2_115.csv')
    document.add_heading('Conversão BCD e redução lógica', level=1)
    add_text(document, 'A conversão binária usa limiares T10, T20 e T30 para selecionar a correção K em 0, 10, 20 ou 30. A unidade é MAG−K por somador/subtrator de cinco bits; no domínio 0..31, R[4] é zero e UNI recebe R[3:0]. O decodificador BCD usa tabela literal ativa em zero para gfedcba; códigos 10..15 apagam os segmentos e foram incluídos na minimização sem dont-cares.')
    tables = root / 'docs' / 'tabelas'
    add_text(document, 'Tabela binário para BCD, cobrindo todas as 32 magnitudes:')
    add_csv_table(document, tables / 'bin_bcd_tabela.csv', ['MAG', 'Binário', 'DEZ', 'UNI'], lambda row: [row['MAG_decimal'], row['MAG_binario'], row['DEZ'], row['UNI']])
    add_text(document, 'Tabela do decodificador, incluindo os seis códigos apagados:')
    add_csv_table(document, tables / 'bcd_7seg_tabela.csv', ['BCD', 'BCD binário', 'SEG gfedcba'], lambda row: [row['BCD_decimal'], row['BCD'], row['SEG_gfedcba']])
    display_path = tables / 'display_decimal_2digitos_tabela.csv'
    boundaries = {0, 9, 10, 19, 20, 29, 30, 31}
    add_text(document, 'Casos representativos dos dois displays com ENABLE ligado e desligado:')
    add_csv_table(document, display_path, ['MAG', 'ENABLE', 'DEZ', 'UNI'], lambda row: [row['MAG'], row['ENABLE'], row['DEZ'], row['UNI']] if int(row['MAG']) in boundaries else None)
    for kmap, caption in (('bin_bcd_kmaps.svg', 'Mapas de Karnaugh dos limiares bin_bcd'), ('bcd_7seg_kmaps.svg', 'Mapas de Karnaugh das sete saídas BCD')):
        source = tables / kmap
        add_picture(document, rasterize_svg(node, node_modules, source), caption, max_height=4.5)
    add_markdown(document, (tables / 'reducoes_bcd_displays.md').read_text(encoding='utf-8'))
    document.add_heading('Tabelas e reduções aritméticas e de controle', level=1)
    arithmetic = tables / 'aritmetica'
    add_text(document, 'Somador completo de um bit: tabela-verdade literal e mapas de Karnaugh que sustentam as equações implementadas.')
    add_csv_table(document, arithmetic / 'somador_1bit.csv', ['A', 'B', 'Cin', 'P', 'S', 'G', 'H', 'Cout'], lambda row: [row['A'], row['B'], row['Cin'], row['P=A XOR B'], row['S'], row['G=A AND B'], row['H=P AND Cin'], row['Cout']])
    add_picture(document, rasterize_svg(node, node_modules, arithmetic / 'kmap_somador_1bit.svg'), 'Mapas de Karnaugh de soma e carry do somador de um bit', max_height=4.5)
    arithmetic_tables = (
        ('sm_para_c2.csv', ['SM[4:0]', 'Sinal', 'Magnitude', 'Valor decimal', 'C2[5:0]'], lambda row: list(row.values())),
        ('c2_para_sm.csv', ['R[5:0]', 'Valor C2', 'Magnitude efetiva[4:0]', 'Sinal efetivo', 'F_SM[5:0]', 'Domínio', 'Observação'], lambda row: list(row.values())),
        ('negador_c2_6bit.csv', ['B_C2[5:0]', 'Valor C2', 'NEG[5:0]', 'Valor C2 de saída'], lambda row: list(row.values())),
    )
    for filename, headers, selector in arithmetic_tables:
        add_text(document, f'Tabela completa: {filename.replace("_", " ").replace(".csv", " ")}.')
        add_csv_table(document, arithmetic / filename, headers, selector)
    for figure in ('cadeia_sm_para_c2.svg', 'cadeia_c2_para_sm.svg', 'cadeia_negador_c2_6bit.svg', 'cadeia_somador_subtrator_5bit.svg', 'cadeia_somador_subtrator_6bit.svg'):
        add_picture(document, rasterize_svg(node, node_modules, arithmetic / figure), f'Redução estrutural — {figure.removesuffix(".svg").replace("_", " ")}', max_height=4.2)
    document.add_heading('Controle e operações bit a bit', level=2)
    add_markdown(document, (tables / 'mapa_logica_5bit.md').read_text(encoding='utf-8-sig'))
    add_markdown(document, (tables / 'tabela_decodificador_operacao.md').read_text(encoding='utf-8-sig'))
    document.add_heading('Circuitos e waveforms dos módulos', level=1)
    for module in MODULES:
        document.add_page_break()
        document.add_heading(module, level=1)
        add_table(document, ['Porta', 'Direção', 'Largura'], interface_rows(interfaces[module]), font_size=8)
        module_text = (root / 'modulos' / f'{module}.bdf.md').read_text(encoding='utf-8')
        module_text = re.sub(r'^# .*\n', '', module_text, count=1)
        add_markdown(document, module_text)
        circuit = circuits[module]
        circuit_image = rasterize_svg(node, node_modules, circuit) if circuit.suffix.lower() == '.svg' else circuit
        add_circuit_panels(document, circuit_image, f'Circuito estrutural de {module}')
        document.add_page_break()
        document.add_heading(f'Waveform de {module}', level=2)
        symbol_text = (root / 'modulos' / f'{module}.bsf.md').read_text(encoding='utf-8')
        symbol_text = re.sub(r'^# .*\n', '', symbol_text, count=1)
        add_markdown(document, symbol_text)
        add_picture(document, waveform_dir / f'{module}.png', f'Waveform simulada de {module}', max_height=5.25)
    document.add_page_break()
    document.add_heading('Integração e resultados da validação', level=1)
    add_text(document, 'A waveform integrada é evidência da execução no modelo exportado da hierarquia. Atribuições de pinos, cobertura de compilação, resultados dos testbenches e pendências são informados conforme VALIDACAO.md.')
    add_picture(document, integration, 'Waveform integrada da ula_de2_115', max_height=5.0)
    add_markdown(document, validation.read_text(encoding='utf-8-sig'))
    document.add_heading('Conclusão', level=1)
    add_text(document, 'A compilação completa da revisão final passou no Quartus com zero erros e 36 avisos, ocupando 100 elementos lógicos e 96 pinos, sem registradores, memória, DSPs ou PLLs. A netlist funcional nativa passou 16.384 verificações sobre 8.192 entradas únicas, e o candidato ZIP extraído recompilou com código de saída zero. Os 16 módulos e a integração também passaram seus testbenches, conforme as contagens registradas. Não houve teste físico na placa.')
    document.save(output)
    print(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--delivery', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--circuit-dir', type=Path, default=Path('entrega/docs/diagramas'))
    parser.add_argument('--waveform-dir', type=Path, default=Path('entrega/simulation/waveforms'))
    parser.add_argument('--validation', type=Path, default=Path('entrega/docs/VALIDACAO.md'))
    parser.add_argument('--output', type=Path, default=Path('entrega/relatorio/RELATORIO.docx'))
    parser.add_argument('--node', type=Path)
    parser.add_argument('--node-modules', type=Path)
    parser.add_argument('--date', default='30/09/2026')
    build(parser.parse_args())


if __name__ == '__main__':
    main()
